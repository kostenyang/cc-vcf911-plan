# -*- coding: utf-8 -*-
import sys
import os
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from content import STEPS, ORIG_REVIEW, GATES, OPEN_ITEMS, DOCS, KBS, IPDNS, GIT_REFS, CAT_ZH, PHASES, LAB_COVERAGE
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "out", "CTBC_VCF911_每步Checklist.xlsx")
FONT = "Arial"
NAVY, PLUM, TEAL, OCEAN, SKY, GREEN, GREY, CRIMSON, GOLD = "1B1D36", "6C4B94", "007B8C", "005C8A", "0098C7", "61A60E", "717074", "A6192E", "F3BA16"
PHASE_COLOR = {"P0": "1D428A", "P1": PLUM, "P2": TEAL, "P3": SKY, "P4": "316E60"}
CAT_COLOR = {"PRE": OCEAN, "EXE": TEAL, "VER": "4E8A0B", "RB": GREY}
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="D0D4D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def f(bold=False, color="222222", size=10, italic=False, underline=None):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic, underline=underline)

def hdr_row(ws, row, headers, fill=NAVY, height=22):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = f(True, "FFFFFF", 10)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = height

def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = f(True, NAVY, 14)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f(False, GREY, 9)

def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def body(cell, wrap=True, bold=False, color="222222", size=10):
    cell.font = f(bold, color, size)
    cell.alignment = Alignment(vertical="top", wrap_text=wrap)
    cell.border = BORDER

wb = Workbook()

# ================================================================ 說明
ws = wb.active
ws.title = "說明"
title(ws, "中信 VCF 9.1.1 Converge & Import — 每一步 Checklist", "版本 v1｜2026-10-07｜Broadcom CXS｜對應簡報：CTBC_VCF911_Converge_Import_Plan.pptx")
rows = [
    ("範圍", "Management Cluster Converge → VCF 9.1.1 管理網域；vRA 8.18 → VCF Automation 9.1.1；整合雲 / 中信雲 Import 為 VI Workload Domain 並升級 9.1.1；SRM 重新註冊；vRO → 9.1.1。不含 NSX Edge cluster。"),
    ("路徑假設", "WLD 採方案 B：vCenter 先以 RDU 升 9.1.1 → SRM Reconfigure → Import（共用管理網域 NSX 9.1.1）→ ESXi 走 VCF LCM。若改方案 A，S07–S14 需改寫（見簡報 Decision 頁）。"),
    ("已確認事項", "2026-10-07：三套 vCenter 皆無 Enhanced Linked Mode（客戶確認）。"),
    ("怎麼填", "只需填黃底欄位：『Checklist』頁的 狀態 / 執行人 / 完成時間 / 證據 / 備註；『Gate 簽核』、『IP & DNS 規劃』、『待確認事項』的黃底欄。其餘為計畫內容，請勿改動。"),
    ("填寫範例", "狀態＝完成｜執行人＝王小明（中信 系統）｜完成時間＝2026-10-19 14:30｜證據＝S04-12_converge-success.png｜備註＝Validation 2 個 Warning 已確認"),
    ("狀態選項", "未開始 / 進行中 / 完成 / 失敗 / N/A（下拉選單；『進度總覽』頁自動統計）"),
    ("類別", "前置檢查 → 執行 → 完成驗證 → 退版 / 截止點。每一步的『退版』列須在 Lab 演練過才可進正式環境。"),
    ("依據欄", "TechDocs / KB = Broadcom 官方文件；CXS Lab 實測 = Broadcom CXS 內部 lab 實測結果（非官方文件）；推定 = 目前無官方文件直接說明，需 Lab 驗證或開 SR 確認。"),
    ("帳密原則", "本表不記錄任何帳號密碼；需要密碼的步驟一律由客戶密碼保管工具提供。"),
    ("負責代號", "PM 專案經理、ARC 架構師、ENG 升級工程師、SYS 系統、HW 硬體、NET 網路、SEC 資安、DR DR 負責人、AUTO 自動化平台、APP 應用驗證、CHG 變更管理、GSS 原廠支援（同人員表）"),
]
r = 4
for k, v in rows:
    ws.cell(row=r, column=1, value=k).font = f(True, NAVY)
    c = ws.cell(row=r, column=2, value=v); c.font = f(); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=r, column=1).alignment = Alignment(vertical="top")
    ws.row_dimensions[r].height = 32 if len(v) > 70 else 18
    r += 1
r += 1
ws.cell(row=r, column=1, value="Gate 與 PONR").font = f(True, NAVY, 12); r += 1
hdr_row(ws, r, ["Gate", "時點 / 說明"]); r += 1
for g, when, name, cond, who, ponr in GATES:
    ws.cell(row=r, column=1, value=g + (" ⚠ PONR" if ponr else "")).font = f(True, CRIMSON if ponr else NAVY)
    c = ws.cell(row=r, column=2, value=f"{when}｜{name}｜放行條件：{cond}｜簽核：{who}")
    c.font = f(); c.alignment = Alignment(wrap_text=True)
    r += 1
widths(ws, [16, 120])
ws.sheet_view.showGridLines = False

# ================================================================ Checklist
wc = wb.create_sheet("Checklist")
title(wc, "每一步 Checklist（方案 B）", "黃底欄位由執行人填寫；狀態請用下拉選單")
H = ["Step", "原計畫", "階段", "類別", "#", "檢查項目", "指令 / UI 路徑", "預期結果", "依據", "負責", "狀態", "執行人", "完成時間", "證據 / 截圖", "備註"]
hdr_row(wc, 3, H)
widths(wc, [7, 12, 9, 10, 8, 52, 38, 18, 34, 9, 10, 12, 16, 22, 22])
row = 4
first_item_row = row
for st in STEPS:
    col = PHASE_COLOR[st["phase"]]
    fill = PatternFill("solid", fgColor=col)
    pz, pe = PHASES[st["phase"]]
    gate = (st["gate"] + (" ⚠ PONR" if st["ponr"] else " Gate")) if st["gate"] else ""
    head = f"{st['id']}｜{st['en']}｜{st['zh']}" + (f"｜{gate}" if gate else "") + (f"｜{st['dur']}" if st["dur"] else "") + f"｜負責：{st['own']}"
    vals = [st["id"], st["orig"], pz, "步驟", "", head]
    for c in range(1, len(H) + 1):
        cell = wc.cell(row=row, column=c, value=vals[c - 1] if c <= len(vals) else None)
        cell.fill = fill
        cell.font = f(True, "FFFFFF", 10.5 if c == 6 else 9)
        cell.alignment = Alignment(vertical="center", wrap_text=(c == 6))
        cell.border = BORDER
    wc.row_dimensions[row].height = 46
    row += 1
    n = 0
    for cat in ["PRE", "EXE", "VER", "RB"]:
        for it in [i for i in st["items"] if i["cat"] == cat]:
            n += 1
            v = [st["id"], st["orig"], pz, CAT_ZH[cat], f"{st['id']}-{n:02d}", it["t"], it["how"], it["exp"], it["ref"], it["own"], "未開始", None, None, None, None]
            for c, val in enumerate(v, 1):
                cell = wc.cell(row=row, column=c, value=val)
                body(cell, size=9.5 if c in (7, 9) else 10)
                if c == 4:
                    cell.font = f(True, CAT_COLOR[cat], 10)
                if c == 6 and (("PONR" in it["t"]) or it["t"].startswith("截止點")):
                    cell.font = f(True, CRIMSON, 10)
                if c >= 11:
                    cell.fill = INPUT_FILL
            row += 1
last_item_row = row - 1
dv = DataValidation(type="list", formula1='"未開始,進行中,完成,失敗,N/A"', allow_blank=True)
wc.add_data_validation(dv)
dv.add(f"K{first_item_row}:K{last_item_row}")
rng = f"K{first_item_row}:K{last_item_row}"
wc.conditional_formatting.add(rng, FormulaRule(formula=[f'K{first_item_row}="完成"'], fill=PatternFill("solid", fgColor="D9EAD3"), font=Font(name=FONT, color="2E6B0E", bold=True)))
wc.conditional_formatting.add(rng, FormulaRule(formula=[f'K{first_item_row}="失敗"'], fill=PatternFill("solid", fgColor="F4CCCC"), font=Font(name=FONT, color=CRIMSON, bold=True)))
wc.conditional_formatting.add(rng, FormulaRule(formula=[f'K{first_item_row}="進行中"'], fill=PatternFill("solid", fgColor="FFE599")))
wc.conditional_formatting.add(rng, FormulaRule(formula=[f'K{first_item_row}="N/A"'], fill=PatternFill("solid", fgColor="E7E6E6"), font=Font(name=FONT, color=GREY)))
wc.freeze_panes = "F4"
wc.auto_filter.ref = f"A3:O{last_item_row}"
wc.print_title_rows = "3:3"
wc.page_setup.orientation = "landscape"
wc.page_setup.fitToWidth = 1
wc.page_setup.fitToHeight = 0
wc.sheet_properties.pageSetUpPr.fitToPage = True

# ================================================================ 進度總覽
wp = wb.create_sheet("進度總覽", 1)
title(wp, "進度總覽（自動統計自『Checklist』頁，請勿手改）")
hdr_row(wp, 3, ["Step", "步驟", "階段", "Gate", "項目數", "完成", "N/A", "失敗", "進行中", "完成率"])
widths(wp, [8, 62, 10, 14, 9, 8, 8, 8, 9, 10])
CL = "Checklist"
rr = 4
for st in STEPS:
    gate = (st["gate"] + (" ⚠ PONR" if st["ponr"] else "")) if st["gate"] else ""
    vals = [st["id"], f"{st['zh']}（{st['orig']}）", PHASES[st["phase"]][0], gate]
    for c, v in enumerate(vals, 1):
        cell = wp.cell(row=rr, column=c, value=v); body(cell)
    A = f"{CL}!$A${first_item_row}:$A${last_item_row}"
    D = f"{CL}!$D${first_item_row}:$D${last_item_row}"
    K = f"{CL}!$K${first_item_row}:$K${last_item_row}"
    wp.cell(row=rr, column=5, value=f'=COUNTIFS({A},$A{rr},{D},"<>步驟")')
    wp.cell(row=rr, column=6, value=f'=COUNTIFS({A},$A{rr},{K},"完成")')
    wp.cell(row=rr, column=7, value=f'=COUNTIFS({A},$A{rr},{K},"N/A")')
    wp.cell(row=rr, column=8, value=f'=COUNTIFS({A},$A{rr},{K},"失敗")')
    wp.cell(row=rr, column=9, value=f'=COUNTIFS({A},$A{rr},{K},"進行中")')
    wp.cell(row=rr, column=10, value=f'=IF(E{rr}=0,0,(F{rr}+G{rr})/E{rr})')
    for c in range(5, 11):
        body(wp.cell(row=rr, column=c))
    wp.cell(row=rr, column=10).number_format = "0%"
    if st["gate"]:
        wp.cell(row=rr, column=4).font = f(True, CRIMSON if st["ponr"] else "7F6000")
    rr += 1
wp.cell(row=rr, column=1, value="合計").font = f(True, NAVY)
for c, col in zip(range(5, 10), "EFGHI"):
    cell = wp.cell(row=rr, column=c, value=f"=SUM({col}4:{col}{rr-1})"); body(cell, bold=True)
cell = wp.cell(row=rr, column=10, value=f"=IF(E{rr}=0,0,(F{rr}+G{rr})/E{rr})"); body(cell, bold=True); cell.number_format = "0%"
wp.cell(row=rr + 2, column=1, value="完成率 =（完成 + N/A）÷ 項目數；『失敗』不算完成，須依該步『退版』列處理後重做。").font = f(False, GREY, 9)
wp.freeze_panes = "A4"

# ================================================================ 原計畫對照
wo = wb.create_sheet("原計畫對照")
title(wo, "原暫定計畫 13 步 → 建議步驟對照", "判定：照做 / 調整 / 移到最後 / 建議架 / 拆兩段")
hdr_row(wo, 3, ["原 #", "原步驟", "判定", "說明", "依據", "建議步驟"])
widths(wo, [6, 40, 11, 70, 40, 12])
VCOL = {"照做": "2E6B0E", "調整": "B45F06", "移到最後": CRIMSON, "建議架": TEAL, "拆兩段": "B45F06"}
for i, rowv in enumerate(ORIG_REVIEW, 4):
    for c, v in enumerate(rowv, 1):
        cell = wo.cell(row=i, column=c, value=v); body(cell)
        if c == 3: cell.font = f(True, VCOL.get(v, "222222"))
wo.freeze_panes = "A4"

# ================================================================ Gate 簽核
wg = wb.create_sheet("Gate 簽核")
title(wg, "Gate 與 Point of No Return 簽核", "PONR = 越過後無官方反向程序；需觀察期與客戶簽核")
hdr_row(wg, 3, ["Gate", "時點", "內容", "放行條件", "簽核角色", "PONR", "結果（Go / No-Go）", "簽核人", "日期", "備註"])
widths(wg, [7, 9, 18, 60, 14, 8, 16, 14, 12, 24])
for i, (g, when, name, cond, who, ponr) in enumerate(GATES, 4):
    vals = [g, when, name, cond, who, "是" if ponr else "否", None, None, None, None]
    for c, v in enumerate(vals, 1):
        cell = wg.cell(row=i, column=c, value=v); body(cell)
        if c == 6 and ponr: cell.font = f(True, CRIMSON)
        if c >= 7: cell.fill = INPUT_FILL
gdv = DataValidation(type="list", formula1='"Go,No-Go,條件式 Go"', allow_blank=True)
wg.add_data_validation(gdv); gdv.add(f"G4:G{3+len(GATES)}")

# ================================================================ IP & DNS
wi = wb.create_sheet("IP & DNS 規劃")
title(wi, "新增元件 IP / DNS 規劃", "黃底欄位填實際值；FQDN 全小寫、A / PTR 都要建（KB 453612、KB 440630）")
hdr_row(wi, 3, ["元件", "IP 數量", "FQDN 需求", "網段", "備註", "實際 FQDN", "實際 IP / 範圍", "VLAN", "A / PTR 驗證", "負責"])
widths(wi, [34, 20, 24, 13, 42, 30, 22, 8, 12, 8])
for i, rowv in enumerate(IPDNS, 4):
    vals = list(rowv) + [None, None, None, None, "NET"]
    for c, v in enumerate(vals, 1):
        cell = wi.cell(row=i, column=c, value=v); body(cell)
        if 6 <= c <= 9: cell.fill = INPUT_FILL
idv = DataValidation(type="list", formula1='"未驗證,正反解 OK,有誤"', allow_blank=True)
wi.add_data_validation(idv); idv.add(f"I4:I{3+len(IPDNS)}")

# ================================================================ 待確認事項
wq = wb.create_sheet("待確認事項")
title(wq, "待確認事項（G0 前關閉）")
hdr_row(wq, 3, ["#", "項目", "為什麼重要", "確認方式", "負責", "狀態", "結論", "日期"])
widths(wq, [5, 30, 64, 18, 10, 10, 40, 12])
items = [("Enhanced Linked Mode", "converge 與 import 都不支援 ELM", "客戶確認", "SYS")] + list(OPEN_ITEMS)
for i, rowv in enumerate(items, 4):
    done = rowv[0] == "Enhanced Linked Mode"
    vals = [i - 3] + list(rowv) + (["已確認", "三套 vCenter 皆無 ELM", "2026-10-07"] if done else ["待確認", None, None])
    for c, v in enumerate(vals, 1):
        cell = wq.cell(row=i, column=c, value=v); body(cell)
        if c >= 6: cell.fill = INPUT_FILL
qdv = DataValidation(type="list", formula1='"待確認,已確認,不適用"', allow_blank=True)
wq.add_data_validation(qdv); qdv.add(f"F4:F{3+len(items)}")
wq.conditional_formatting.add(f"F4:F{3+len(items)}", FormulaRule(formula=['F4="已確認"'], fill=PatternFill("solid", fgColor="D9EAD3")))

# ================================================================ 官方文件
wd = wb.create_sheet("官方文件")
title(wd, "官方依據：Broadcom TechDocs 與 KB")
hdr_row(wd, 3, ["類別", "文件", "URL"])
widths(wd, [12, 70, 90])
r = 4
for cat, t, u in DOCS:
    body(wd.cell(row=r, column=1, value=cat)); body(wd.cell(row=r, column=2, value=t))
    c = wd.cell(row=r, column=3, value=u); c.hyperlink = u; body(c, wrap=False); c.font = f(False, "0563C1", 9, underline="single")
    r += 1
for n, t in KBS:
    u = f"https://knowledge.broadcom.com/external/article/{n}"
    body(wd.cell(row=r, column=1, value=f"KB {n}")); body(wd.cell(row=r, column=2, value=t))
    c = wd.cell(row=r, column=3, value=u); c.hyperlink = u; body(c, wrap=False); c.font = f(False, "0563C1", 9, underline="single")
    r += 1

# ================================================================ Lab 驗證覆蓋
wl = wb.create_sheet("Lab 驗證覆蓋")
title(wl, "CXS Lab 驗證覆蓋（github.com/kostenyang）", "未驗證的步驟正式變更前須於 Lab 補測")
hdr_row(wl, 3, ["步驟", "內容", "狀態", "實測日期", "結果", "差異 / 缺口", "Git 出處"])
widths(wl, [10, 30, 9, 22, 60, 50, 44])
for i, rowv in enumerate(LAB_COVERAGE, 4):
    for c, v in enumerate(rowv, 1):
        cell = wl.cell(row=i, column=c, value=v); body(cell)
        if c == 3: cell.font = f(True, "2E6B0E" if v == "已驗證" else CRIMSON)

# ================================================================ Git 參考文件
wr = wb.create_sheet("Git 參考文件")
title(wr, "Git 參考文件（github.com/kostenyang）— 內部參考", "CXS Lab 實測紀錄，非官方文件；private repo 內含 lab 資訊，交付客戶前可刪除本頁")
hdr_row(wr, 3, ["對應步驟", "Repo", "路徑", "可參考的內容"])
widths(wr, [11, 30, 58, 80])
for i, (stp, repo, path, what) in enumerate(GIT_REFS, 4):
    vals = [stp, repo, path, what]
    for c, v in enumerate(vals, 1):
        cell = wr.cell(row=i, column=c, value=v); body(cell)
    name = repo.split("（")[0].split(" /")[0].strip()
    if name and name != "—":
        cell = wr.cell(row=i, column=2)
        cell.hyperlink = f"https://github.com/kostenyang/{name}"
        cell.font = f(False, "0563C1", 10, underline="single")

for sh in wb.worksheets:
    sh.sheet_view.zoomScale = 100
    sh.page_setup.orientation = "landscape"
    sh.page_setup.paperSize = sh.PAPERSIZE_A4
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
    sh.sheet_properties.pageSetUpPr.fitToPage = True
    sh.page_margins.left = sh.page_margins.right = 0.4
wb.save(OUT)
print("saved", OUT, "items rows", first_item_row, last_item_row)
