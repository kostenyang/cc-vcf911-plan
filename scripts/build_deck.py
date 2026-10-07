# -*- coding: utf-8 -*-
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

import os
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from content import STEPS, STEP_BY_ID, ORIG_REVIEW, GATES, OPEN_ITEMS, DOCS, KBS, IPDNS, CAT_ZH, LAB_COVERAGE

BASE = os.environ.get("BASE_PPTX", os.path.join(HERE, "..", "template", "TECH_TUESDAY_Whats_New_with_vSAN_in_VCF_9_1.pptx"))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "out", "CTBC_VCF911_Converge_Import_Plan.pptx")

def rgb(h): return RGBColor.from_string(h)
PLUM, TEAL, OCEAN, SKY, GREEN, GOLD = "6C4B94", "007B8C", "005C8A", "0098C7", "61A60E", "F3BA16"
NAVY, GREY, LIGHT, PALE, CRIMSON, ORANGE, ROYAL, DTEAL, DARK = "1B1D36", "717074", "F2F2F2", "E8F0F4", "A6192E", "E68C28", "1D428A", "316E60", "333333"
PHASE_COLOR = {"P0": ROYAL, "P1": PLUM, "P2": TEAL, "P3": SKY, "P4": DTEAL}
CAT_COLOR = {"PRE": OCEAN, "EXE": TEAL, "VER": GREEN, "RB": GREY}
CAT_EN = {"PRE": "Pre-check", "EXE": "Execute", "VER": "Verify", "RB": "Rollback"}
EA = "Microsoft JhengHei"

prs = Presentation(BASE)
M0 = prs.slide_masters[0]
def layout(name):
    for l in M0.slide_layouts:
        if l.name == name: return l
    raise KeyError(name)

orig_ids = list(prs.slides._sldIdLst)
COVER_ID, THANKS_ID = orig_ids[0], orig_ids[27]

# ---------------------------------------------------------------- text helpers
def set_run_font(run, size=None, bold=None, color=None, italic=None):
    f = run.font
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color: f.color.rgb = rgb(color)
    f.name = "Arial"
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", EA if tag == "a:ea" else "Arial")

def fill_tf(tf, paras, size=11, color=DARK, bold=False, align=PP_ALIGN.LEFT, space_after=2, line=None):
    """paras: list of str | list of (text, dict) runs | dict(runs=..., bullet=..)"""
    tf.word_wrap = True
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = align
        para.space_after = Pt(space_after)
        if line: para.line_spacing = line
        runs = p if isinstance(p, list) else [(p, {})]
        for text, st in runs:
            r = para.add_run()
            r.text = text
            set_run_font(r, size=st.get("size", size), bold=st.get("bold", bold),
                         color=st.get("color", color), italic=st.get("italic"))
    return tf

def tbox(slide, x, y, w, h, paras, size=11, color=DARK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, margin=0.05, space_after=2, line=None):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    fill_tf(tf, paras, size, color, bold, align, space_after, line)
    return s

def rect(slide, x, y, w, h, fillc=None, linec=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fillc:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fillc)
    else:
        s.fill.background()
    if linec:
        s.line.color.rgb = rgb(linec); s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    return s

def shape_text(s, paras, size=11, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.06):
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    fill_tf(tf, paras, size, color, bold, align, 0)
    return s

def add_slidenum(slide):
    s = slide.shapes.add_textbox(Inches(12.55), Inches(7.05), Inches(0.55), Inches(0.25))
    p = s.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(p._p, qn("a:fld"))
    fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}"); fld.set("type", "slidenum")
    rpr = etree.SubElement(fld, qn("a:rPr")); rpr.set("lang", "en-US"); rpr.set("sz", "800")
    sf = etree.SubElement(rpr, qn("a:solidFill")); c = etree.SubElement(sf, qn("a:srgbClr")); c.set("val", GREY)
    t = etree.SubElement(fld, qn("a:t")); t.text = "‹#›"

def content_slide(title, subtitle, sub_color=TEAL, source=None):
    s = prs.slides.add_slide(layout("Title Only"))
    s.shapes.title.text = title
    for r in s.shapes.title.text_frame.paragraphs[0].runs:
        rPr = r._r.get_or_add_rPr()
        ea = rPr.find(qn("a:ea"))
        if ea is None: ea = etree.SubElement(rPr, qn("a:ea"))
        ea.set("typeface", EA)
    if subtitle:
        tbox(s, 0.6, 0.88, 12.1, 0.32, [subtitle], size=12.5, color=sub_color, bold=True, margin=0.03)
    if source:
        tbox(s, 0.6, 6.66, 12.1, 0.26, [source], size=7.5, color=GREY, margin=0.03)
    add_slidenum(s)
    return s

def section(title, subtitle, color_name):
    s = prs.slides.add_slide(layout(f"Section Header – {color_name}"))
    for ph in s.placeholders:
        idx = ph.placeholder_format.idx
        if idx == 0: ph.text = title
        elif idx == 10: ph.text = subtitle
        for p in ph.text_frame.paragraphs:
            for r in p.runs:
                rPr = r._r.get_or_add_rPr()
                ea = etree.SubElement(rPr, qn("a:ea")); ea.set("typeface", EA)
    return s

def pill(slide, x, y, w, h, text, fillc, size=9, color="FFFFFF"):
    s = rect(slide, x, y, w, h, fillc, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    s.adjustments[0] = 0.5
    shape_text(s, [text], size=size, color=color, bold=True, margin=0.03)
    return s

def table(slide, x, y, w, col_w, rows, header_fill=NAVY, size=9.5, row_h=0.3, hdr_size=10, fills=None, bold_cols=()):
    nrows, ncols = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
    tblPr = gt._element.graphic.graphicData.tbl.tblPr
    # remove built-in style banding look
    tblPr.set("bandRow", "0"); tblPr.set("firstRow", "1")
    tbl = gt.table
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for r in range(nrows):
        tbl.rows[r].height = Inches(row_h)
        for c in range(ncols):
            cell = tbl.cell(r, c)
            val = rows[r][c]
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.025)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.paragraphs[0].text = ""
            if r == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb(header_fill)
                fill_tf(tf, [val], hdr_size, "FFFFFF", True, PP_ALIGN.LEFT, 0)
            else:
                f = None
                if fills: f = fills(r, c, val)
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb(f[0] if f else ("FFFFFF" if r % 2 else "F5F7F9"))
                tcolor = f[1] if f else DARK
                fill_tf(tf, [val], size, tcolor, (c in bold_cols) or bool(f and f[2]), PP_ALIGN.LEFT, 0)
    return gt

def build_src(reflist, maxlen=175):
    td, kb, other = [], [], []
    for ref in reflist:
        for part in ref.replace("；", ";").split(";"):
            p = part.strip()
            if not p: continue
            if p.startswith("TechDocs:"):
                t = p.replace("TechDocs:", "").strip().split("（")[0].strip()
                if not any(t.startswith(x) or x.startswith(t) for x in td): td.append(t)
            elif p.startswith("KB "):
                n = p[3:].strip()
                if n not in kb: kb.append(n)
            else:
                q = "CXS Lab 實測" if p.startswith("CXS Lab") else p
                if q not in other: other.append(q)
    def fmt(tdl, more):
        segs = []
        if tdl: segs.append("TechDocs" + "".join(f"《{t}》" for t in tdl) + ("等" if more else ""))
        if kb: segs.append("KB " + "、".join(kb))
        segs += other
        return "依據：" + "；".join(segs)
    out = fmt(td, False)
    k = len(td)
    while len(out) > maxlen and k > 1:
        k -= 1
        out = fmt(td[:k], True)
    return out if (td or kb or other) else None

def short_ref(ref):
    out = []
    for part in ref.replace("；", ";").split(";"):
        p = part.strip()
        if not p: continue
        if p.startswith("CXS Lab") or p == "推定": p = "CXS Lab 實測" if p.startswith("CXS") else "推定"
        p = p.replace("TechDocs: ", "TechDocs《").strip()
        if p.startswith("TechDocs《"): p = p + "》"
        out.append(p)
    return out

# ================================================================= slides
new_slides = []

# ---- 2 Plan at a glance
s = content_slide("Plan at a Glance", "結論先講：原 13 步大方向可行，4 個地方要調整", TEAL)
cols = [
    ("01", PLUM, "Sequence / 順序調整", ["vRO 移到最後（VCFA 9.1.1 之後）", "vRA / vRO 先冷遷移進管理網域", "WLD 升級拆成 vCenter、ESXi 兩段"]),
    ("02", TEAL, "WLD Path / 匯入路徑", ["import 8.0.3 一定會帶 NSX（新 3-node）", "建議：vCenter 先 RDU → import 共用 NSX", "設計會議定案 A / B"]),
    ("03", SKY, "Depot / 離線來源", ["Offline Depot 非必要，但建議架", "Installer 內 binaries 不會轉移", "VCFA、WLD 升級都還要用"]),
    ("04", GREEN, "Gates & PONR / 管控", ["G0–G5 六個 Gate", "PONR：Converge、兩次 Import、vSAN format", "每一步都有退版點與截止點"]),
]
x0, cw = 0.63, 2.95
for i, (num, col, head, bl) in enumerate(cols):
    x = x0 + i * (cw + 0.12)
    c = rect(s, x + 0.05, 1.45, 0.8, 0.8, col, shape=MSO_SHAPE.OVAL)
    shape_text(c, [num], size=18, bold=True)
    if i < 3:
        ln = s.shapes.add_connector(1, Inches(x + 0.95), Inches(1.85), Inches(x + cw + 0.1), Inches(1.85))
        ln.line.color.rgb = rgb("BFBFBF"); ln.line.width = Pt(1)
    tbox(s, x, 2.38, cw, 0.36, [head], size=13, color=col, bold=True)
    bar = rect(s, x + 0.04, 2.78, cw - 0.12, 0.04, col)
    tbox(s, x, 2.95, cw, 1.8, ["• " + b for b in bl], size=12.5, color=DARK, space_after=8)
box = rect(s, 0.63, 4.95, 12.07, 0.95, "F5F7F9", "D9D9D9", shape=MSO_SHAPE.ROUNDED_RECTANGLE)
box.adjustments[0] = 0.08
tbox(s, 0.8, 5.02, 11.8, 0.85, [
    [("Scope ｜ ", {"bold": True, "color": NAVY}), ("Management Cluster Converge · vRA 8.18 → VCF Automation 9.1.1 · 整合雲 / 中信雲 Import 與升級 · SRM 重新註冊 · vRO → 9.1.1", {"color": DARK})],
    [("前提 ｜ ", {"bold": True, "color": NAVY}), ("不含 NSX Edge cluster（無 Edge node / T0 / T1）；每一步都要有已在 Lab 演練過的退版點，才進正式環境。", {"color": DARK})],
], size=11, space_after=4, anchor=MSO_ANCHOR.MIDDLE)
new_slides.append(s)

# ---- 3 section
new_slides.append(section("Current & Target", "現況、目標架構與原計畫檢視", "Plum"))

# ---- 4 architecture
s = content_slide("Current vs Target ｜ One VCF 9.1.1 Instance",
                  "一個 VCF Instance：管理網域 + 兩個 VI Workload Domain", PLUM,
                  "依據：TechDocs《Supported and Not Supported Configurations to Converge to VCF》《Import an Existing vCenter to Create a Workload Domain》《Upgrading to VCF Automation 9.1》")
tbox(s, 0.63, 1.3, 4.5, 0.3, ["Current ｜ 現況"], size=11, color=GREY, bold=True)
tbox(s, 5.85, 1.3, 6.8, 0.3, ["Target ｜ VCF 9.1.1 Instance"], size=11, color=PLUM, bold=True)
rows = [
    ("MANAGEMENT", PLUM,
     ["Management Cluster", "vCenter 9.1.1 + ESXi 9.1.1 ×6（vSphere，尚非 VCF）"],
     ["VCF 9.1.1 Management Domain（Converge）", "SDDC Manager · NSX Manager（無 Edge）· VCF Operations + Mgmt Services · Identity Broker · License Server",
      "VCF Automation 9.1.1（vRA 8.18 藍綠升級）· VCF Operations orchestrator 9.1.1（vRO）"]),
    ("WORKLOAD DOMAIN ｜ 整合雲", TEAL,
     ["整合雲 Cluster", "vCenter 8.0.3 + ESXi 8.0.3 ×6 · SRM 9.1.1"],
     ["VI Workload Domain（Import）", "vCenter 9.1.1 + ESXi 9.1.1 ×6 · 共用管理網域 NSX 9.1.1（方案 B）",
      "SRM 9.1.1：vCenter 升級後兩站 Reconfigure + Reconnect"]),
    ("WORKLOAD DOMAIN ｜ 中信雲", SKY,
     ["中信雲 Cluster", "vCenter 8.0.3 + ESXi 8.0.3 ×6 · SRM 9.1.1", "vRA 8.18 · vRO（將移出至管理網域）"],
     ["VI Workload Domain（Import）", "vCenter 9.1.1 + ESXi 9.1.1 ×6 · 共用管理網域 NSX 9.1.1（方案 B）",
      "SRM 9.1.1 重新註冊；vRA / vRO 已不在此叢集"]),
]
y = 1.65
for tag, col, cur, tgt in rows:
    tbox(s, 0.63, y, 4.5, 0.22, [tag], size=8, color=col, bold=True, margin=0.02)
    cb = rect(s, 0.63, y + 0.24, 4.6, 1.2, "FFFFFF", col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.25); cb.adjustments[0] = 0.06
    tbox(s, 0.75, y + 0.3, 4.4, 1.1, [[(cur[0], {"bold": True, "color": col, "size": 11.5})]] + [c for c in cur[1:]],
         size=10, color=DARK, space_after=3)
    ar = rect(s, 5.32, y + 0.66, 0.42, 0.36, col, shape=MSO_SHAPE.RIGHT_ARROW)
    tb = rect(s, 5.85, y + 0.24, 6.85, 1.2, col, shape=MSO_SHAPE.ROUNDED_RECTANGLE); tb.adjustments[0] = 0.06
    tbox(s, 5.97, y + 0.3, 6.65, 1.1, [[(tgt[0], {"bold": True, "size": 11.5})]] + [t for t in tgt[1:]],
         size=10, color="FFFFFF", space_after=3)
    y += 1.62
new_slides.append(s)

# ---- 5 original plan review
s = content_slide("Original Plan Review ｜ 13 Steps Checked against TechDocs",
                  "原暫定計畫逐步檢視：照做 / 調整 / 移位，以及對應的建議步驟", PLUM,
                  "依據：TechDocs VCF 9.1（Converge、Import vCenter、Import / Upgrade Aria Automation、Upgrade VCF Operations Orchestrator、Downloading Binaries）；KB 318582、446701；Broadcom Interop Matrix")
hdr = ["#", "原步驟", "判定", "說明", "建議步驟"]
rows_t = [hdr] + [[a, b, c, d, f] for (a, b, c, d, e, f) in ORIG_REVIEW]
VC = {"照做": (GREEN, True), "調整": (ORANGE, True), "移到最後": (CRIMSON, True), "建議架": (TEAL, True), "拆兩段": (ORANGE, True)}
def fills_review(r, c, v):
    if c == 2 and v in VC: return (VC[v][0], "FFFFFF", True)
    return None
table(s, 0.63, 1.32, 12.07, [0.4, 3.1, 0.95, 6.32, 1.3], rows_t, size=9.5, row_h=0.36, fills=fills_review, bold_cols=(4,))
new_slides.append(s)

# ---- 6 decision A/B
s = content_slide("Decision ｜ How the Two vSphere 8 Clouds Enter VCF",
                  "Import 8.0.3 vCenter 一定會帶 NSX → 兩種路徑擇一（建議方案 B）", PLUM,
                  "依據：TechDocs《Import an Existing vCenter to Create a Workload Domain》（NSX 9.1 does not support vCenter 8.0 Update 3a or later）、《Sequential / Optimized Update of vCenter and NSX》；KB 429205；CXS Lab 實測（v8tov9 import-wld）")
def option(x, w, head, col, flow, bullets, rec=False):
    hb = rect(s, x, 1.35, w, 0.45, col); shape_text(hb, [head], size=12.5, bold=True, align=PP_ALIGN.LEFT, margin=0.12)
    if rec: pill(s, x + w - 1.25, 1.42, 1.15, 0.3, "建議 Recommended", GOLD, size=8.5, color=NAVY)
    body = rect(s, x, 1.8, w, 3.55, "F5F7F9", "D9D9D9")
    fx, fw = x + 0.12, (w - 0.24 - 0.09 * (len(flow) - 1)) / len(flow)
    for i, f in enumerate(flow):
        ch = rect(s, fx + i * (fw + 0.09), 1.95, fw, 0.62, "FFFFFF", col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.25)
        ch.adjustments[0] = 0.15
        shape_text(ch, [f], size=9, color=col, bold=True, margin=0.03)
    tbox(s, x + 0.12, 2.72, w - 0.24, 2.6, bl_fmt(bullets), size=11, color=DARK, space_after=6)
def bl_fmt(bullets):
    out = []
    for b in bullets:
        mark, text = b[0], b[1:]
        colr = {"+": GREEN, "−": CRIMSON, "!": ORANGE}[mark]
        out.append([(("✔ " if mark == "+" else "✖ " if mark == "−" else "▲ "), {"color": colr, "bold": True}), (text, {})])
    return out
option(0.63, 5.95, "方案 A ｜ 照原順序：先 Import 8.0.3，再在 VCF 內升級", GREY,
       ["Import\n(+NSX 4.2 ×3)", "NSX\n4.2 → 9.1.1", "vCenter RDU\n(SDDC Mgr)", "ESXi\n9.1.1"],
       ["+vCenter 升級走 SDDC Manager LCM（含 precheck）",
        "−每朵雲新部署 3-node NSX Manager（4.2.x，裝在該叢集）→ 多 6 台 NSX VM",
        "−WLD 升級變三段：NSX → vCenter → ESXi",
        "!Import 的 NSX 版本須在 9.1.1 升級矩陣內（KB 429205 控版）",
        "!SDDC Manager 的 RDU 完成後會刪除來源 vCenter VM"])
option(6.75, 5.95, "方案 B ｜ vCenter 先升 9.1.1，再 Import 共用 NSX", GREEN,
       ["vCenter\nRDU 9.1.1", "SRM\nReconfigure", "Import\n(共用 NSX)", "ESXi 9.1.1\n(VCF LCM)"],
       ["+共用管理網域 NSX 9.1.1：不新增 NSX VM、不必升 NSX",
        "+官方明文：要用 NSX 9.1 須先把 vCenter 升到 9.1",
        "+CXS Lab 已實測：import 57/57 成功、約 15 分、ESXi 8.0.3 照常運作",
        "!vCenter 升級在 VCF 外以 RDU 執行 → 用 VAMI precheck + Lab 演練補強",
        "!Import 前須在管理網域 NSX 註冊 Compute Manager"], rec=True)
cb = rect(s, 0.63, 5.5, 12.07, 0.95, "FFF8E6", GOLD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.25); cb.adjustments[0] = 0.1
tbox(s, 0.8, 5.55, 11.8, 0.85, [
    [("決策需求 ｜ ", {"bold": True, "color": NAVY}), ("G0 前由 ARC + PM 定案。建議方案 B（本簡報與 Excel 的 S07–S14 以方案 B 撰寫）。", {})],
    [("若選方案 A ｜ ", {"bold": True, "color": NAVY}), ("Import 前以 KB 429205 控制 NSX 版本；WLD 升級改為 NSX → vCenter → ESXi，退版改以 file-based restore 為主（來源 vCenter VM 會被刪除）。", {})],
], size=11, color=DARK, space_after=4, anchor=MSO_ANCHOR.MIDDLE)
new_slides.append(s)

# ---- 7 section
new_slides.append(section("Recommended Flow", "建議執行順序、Gate 與 Point of No Return", "Aqua"))

# ---- 8 flow overview
s = content_slide("Recommended Flow ｜ 17 Steps in 5 Phases",
                  "管理網域先做對，再一朵雲一朵雲接進來；方括號內為原計畫步驟編號", TEAL,
                  "紅色 = Point of No Return（需簽核）　金色 = Gate　｜ 依據：TechDocs《Upgrading to VCF 9.1.x》元件順序；KB 440630")
lanes = [
    ("P0", "準備", [("S00", "盤點 / 備份", "新增", "G0", False)]),
    ("P1", "管理網域", [("S01", "vRA / vRO 冷遷移", "#1", "", False), ("S02", "VCF Installer", "#3", "", False),
                       ("S03", "Depot / Binaries", "#4", "", False), ("S04", "Converge", "#5", "G1", True),
                       ("S05", "Import vRA", "#6", "", False), ("S06", "VCFA 9.1.1", "#7", "G2", False)]),
    ("P2", "整合雲", [("S07", "vCenter RDU", "#10", "", False), ("S08", "SRM 重新註冊", "#11", "", False),
                     ("S09", "Import WLD", "#8", "G3", True), ("S10", "ESXi 9.1.1", "#10", "", False)]),
    ("P3", "中信雲", [("S11", "vCenter RDU", "#12", "", False), ("S12", "SRM 重新註冊", "#13", "", False),
                     ("S13", "Import WLD", "#9", "G4", True), ("S14", "ESXi 9.1.1", "#12", "", False)]),
    ("P4", "收尾", [("S15", "vRO 9.1.1", "#2", "", False), ("S16", "vSAN / 收尾", "新增", "G5", True)]),
]
y = 1.38
for pid, pname, steps in lanes:
    col = PHASE_COLOR[pid]
    lb = rect(s, 0.63, y, 1.45, 0.86, col, shape=MSO_SHAPE.ROUNDED_RECTANGLE); lb.adjustments[0] = 0.12
    shape_text(lb, [[(pid + "\n", {"size": 9})], [(pname, {"size": 12})]], size=12)
    bx, bw = 2.2, 1.66
    for i, (sid, name, orig, gate, ponr) in enumerate(steps):
        x = bx + i * (bw + 0.1)
        b = rect(s, x, y, bw, 0.86, "FFFFFF", CRIMSON if ponr else col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=2 if ponr else 1.25)
        b.adjustments[0] = 0.12
        shape_text(b, [[(sid + "  ", {"bold": True, "color": col, "size": 11}), (f"[{orig}]", {"color": GREY, "size": 9, "bold": False})],
                       [(name, {"color": DARK, "size": 12, "bold": True})]], size=11, color=DARK, margin=0.04)
        if gate:
            pill(s, x + bw - 0.78, y - 0.11, 0.76, 0.24, gate + (" PONR" if ponr else ""), CRIMSON if ponr else GOLD,
                 size=8, color="FFFFFF" if ponr else NAVY)
        if i < len(steps) - 1:
            a = rect(s, x + bw + 0.0, y + 0.33, 0.1, 0.2, "9A9A9A", shape=MSO_SHAPE.RIGHT_ARROW)
    y += 1.03
new_slides.append(s)

# ---- 9 rollback summary
s = content_slide("Rollback Means & Cut-off by Step",
                  "每一步「用什麼退」以及「退到哪個時間點為止」", TEAL,
                  "依據：KB 448258、436924、430524、441333、446701、448322、316592；TechDocs《Optimized Update of vCenter and NSX》《About the vSAN Disk Format》；CXS Lab 實測")
rb_rows = [["步驟", "退版手段", "截止點 / PONR"],
    ["S01 vRA / vRO 冷遷移", "反向冷遷移回中信雲；或開回 Clone 保留的來源（保持關機）", "S05 Import 前"],
    ["S04 Converge", "完成前失敗：刪除已部署元件與 Installer，依 JSON 重做（KB 448258）", "⚠ 成功即 PONR（G1）"],
    ["S05 Import vRA", "只是註冊、來源不變；需撤除依 KB 441333 cleanup", "—"],
    ["S06 VCFA 9.1.1", "Cutover 前來源原封不動；Cutover 後開回 8.18 / revert snapshot + 開 SR", "刪除 8.18 來源 VM / snapshot"],
    ["S07 / S11 vCenter RDU", "Switchover 前 RDU 自動回復；之後以 file-based backup restore", "刪除退役來源 VM / 舊備份"],
    ["S08 / S12 SRM 重新註冊", "Reconfigure 不破壞設定；失敗依 KB 448322 修正後重做", "—"],
    ["S09 / S13 Import WLD", "無 UI 反向；KB 436924（API 移除 domain）+ KB 430524", "⚠ Import 完成即 PONR（G3 / G4）"],
    ["S10 / S14 ESXi", "單台 Shift+R 回 altbootbank（KB 316592；vLCM image 影響待 Lab 驗證）", "vSAN on-disk format 升級"],
    ["S15 vRO", "Revert 升級前 snapshot（不含 memory）", "刪除 snapshot"],
    ["S16 vSAN format", "無退版手段", "⚠ 執行即 PONR（G5）"]]
def fills_rb(r, c, v):
    if c == 2 and v.startswith("⚠"): return ("FBE9EB", CRIMSON, True)
    return None
table(s, 0.63, 1.32, 12.07, [2.45, 6.72, 2.9], rb_rows, size=9.5, row_h=0.43, fills=fills_rb, bold_cols=(0,))
new_slides.append(s)

# ---- lab coverage
s = content_slide("Lab Coverage ｜ What Has Been Tested in CXS Lab",
                  "每一步是否已在 CXS Lab（VCF 9.1.1 nested）實測過；未驗證者正式前須補測", TEAL,
                  "來源：github.com/kostenyang（v8tov9、vra-lifecycle、vcf9offlinescript、debug-vcf9.1）— CXS Lab 實測紀錄，非官方文件")
rows_l = [["步驟", "內容", "狀態", "實測日期", "結果 / 差異"]] + [[a, b, c, d, e + ("；" + f if f != "—" else "")] for a, b, c, d, e, f, g in LAB_COVERAGE]
def fills_lab(r, c, v):
    if c == 2:
        return ("D9EAD3", "2E6B0E", True) if v == "已驗證" else ("FBE9EB", CRIMSON, True)
    return None
table(s, 0.63, 1.3, 12.07, [1.05, 2.75, 0.85, 1.9, 5.52], rows_l, size=8.5, row_h=0.42, hdr_size=9.5, fills=fills_lab, bold_cols=(0,))
new_slides.append(s)

# ---- 10 section
new_slides.append(section("Step-by-step Checklist", "每一動的前置檢查、執行、驗證與退版", "Leaf"))

# ---- step slides

import math
def text_w(t, size):
    w = 0.0
    for ch in t:
        w += (size / 72.0) * (1.0 if ord(ch) > 0x2E80 else 0.55)
    return w
def col_h(items, size, width):
    h = 0.0
    for main, how in items:
        lines = math.ceil(text_w("☐ " + main, size) / width)
        h += lines * size * 1.22 / 72.0 + 7 / 72.0
        if how:
            hs = size - 1.5
            h += math.ceil(text_w("    " + how, hs) / width) * hs * 1.22 / 72.0 + 3 / 72.0
    return h
def fit_size(colsdata, width, height):
    for size in [12.5, 12, 11.5, 11, 10.5, 10, 9.5, 9, 8.5, 8]:
        if all(col_h(v, size, width) <= height for v in colsdata.values()):
            return size
    return 8
def step_slide(st):
    col = PHASE_COLOR[st["phase"]]
    src = build_src([it["ref"] for it in st["items"]])
    s = content_slide(f"{st['id']} ｜ {st['en']}", f"{st['zh']}", col, src)
    meta = f"原計畫 {st['orig']} ｜ 負責：{st['own']}" + (f" ｜ {st['dur']}" if st["dur"] else "")
    tbox(s, 0.6, 1.2, 9.5, 0.28, [meta], size=10, color=GREY, margin=0.03)
    if st["gate"]:
        pill(s, 11.3 if st["ponr"] else 11.85, 1.2, 1.4 if st["ponr"] else 0.85, 0.27,
             st["gate"] + (" ⚠ PONR" if st["ponr"] else " Gate"), CRIMSON if st["ponr"] else GOLD, size=9,
             color="FFFFFF" if st["ponr"] else NAVY)
    cw, gap, x0, y0, hh, bh = 2.93, 0.117, 0.63, 1.55, 0.42, 4.55
    colsdata = {}
    for cat in ["PRE", "EXE", "VER", "RB"]:
        colsdata[cat] = [((it["d"] or it["t"]), (it["how"] if st["id"] != "S00" else "")) for it in st["items"] if it["cat"] == cat]
    size = fit_size(colsdata, cw - 0.2, bh - 0.2)
    for i, cat in enumerate(["PRE", "EXE", "VER", "RB"]):
        x = x0 + i * (cw + gap)
        hc = CRIMSON if (cat == "RB" and st["ponr"]) else CAT_COLOR[cat]
        h = rect(s, x, y0, cw, hh, hc)
        shape_text(h, [[(f"{CAT_ZH[cat]}", {"size": 12}), (f"  {CAT_EN[cat]}", {"size": 9, "bold": False})]], size=12, align=PP_ALIGN.LEFT, margin=0.1)
        rect(s, x, y0 + hh, cw, bh, "F5F7F9", "E1E4E8")
        paras = []
        for main, how in colsdata[cat]:
            paras.append([("☐ ", {"color": hc, "bold": True}), (main, {})])
            if how:
                paras.append([("    " + how, {"color": GREY, "size": size - 1.5})])
        tb = tbox(s, x + 0.06, y0 + hh + 0.08, cw - 0.12, bh - 0.12, paras, size=size, color=DARK, space_after=3, line=1.0)
        # extra space after each item (main or how) – set on the last paragraph of each item
        tf = tb.text_frame
        idx = 0
        for main, how in colsdata[cat]:
            last = idx + (1 if how else 0)
            tf.paragraphs[last].space_after = Pt(7)
            idx = last + 1
    return s

for sid in ["S00", "S01", "S02", "S03", "S04", "S05", "S06", "S07", "S08", "S09", "S10"]:
    new_slides.append(step_slide(STEP_BY_ID[sid]))

# combined S11–S14
s = content_slide("S11–S14 ｜ Repeat for 中信雲", "中信雲：步驟同 S07–S10，只列差異", PHASE_COLOR["P3"],
                  "依據：KB 446701（每次 vCenter 升級後兩站 Reconfigure）；TechDocs《Order of Upgrading vSphere and Protection and Recovery Components》《Import an Existing vCenter》")
tbox(s, 0.6, 1.2, 11, 0.28, ["原計畫 #9、#12、#13 ｜ 負責：ENG、SYS、DR、ARC ｜ 詳細 checklist 見 Excel（S11–S14 已逐項展開）"], size=10, color=GREY, margin=0.03)
pill(s, 11.3, 1.2, 1.4, 0.27, "G4 ⚠ PONR", CRIMSON, size=9)
sub = [
    ("S11 vCenter RDU", SKY, ["☐ 同 S07 全部項目", "☐ 若中信雲為 recovery site，排在 protected site 之後", "☐ RDU 暫時 IP 用中信雲網段", "☐ vRA / vRO 已移出；確認 vIDM / ASL 不受 RDU 停機影響"]),
    ("S12 SRM 重新註冊", SKY, ["☐ 同 S08 全部項目", "☐ 兩站 VAMI Reconfigure 再做一次（每次 vCenter 升級後都要）", "☐ Test Recovery 兩個方向都驗"]),
    ("S13 Import WLD", CRIMSON, ["☐ 同 S09 全部項目", "☐ 中信雲 vCenter 先在管理網域 NSX 註冊 Compute Manager", "☐ ⚠ Import 完成即 PONR → G4 簽核"]),
    ("S14 ESXi 9.1.1", SKY, ["☐ 同 S10 全部項目", "☐ 完成後三個叢集皆 9.1.1", "☐ vSAN on-disk format 統一留到 S16"]),
]
for i, (h, c, bl) in enumerate(sub):
    x = 0.63 + i * (2.93 + 0.117)
    hb = rect(s, x, 1.58, 2.93, 0.42, c); shape_text(hb, [h], size=12, align=PP_ALIGN.LEFT, margin=0.1)
    rect(s, x, 2.0, 2.93, 4.48, "F5F7F9", "E1E4E8")
    tbox(s, x + 0.06, 2.08, 2.81, 4.3, bl, size=10.5, color=DARK, space_after=8)
new_slides.append(s)

for sid in ["S15", "S16"]:
    new_slides.append(step_slide(STEP_BY_ID[sid]))

# ---- section
new_slides.append(section("Open Items & References", "待確認事項、IP / DNS 與官方依據", "Plum"))

# ---- open items
s = content_slide("Open Items to Confirm before G0", "以下內容決定實際順序與設計，需在 G0 前確認", PLUM,
                  "標示「推定」者目前無官方文件直接說明，需以 Lab 驗證或 SR 確認")
rows_o = [["待確認項目", "為什麼重要", "確認方式", "負責"]] + [list(r) for r in OPEN_ITEMS]
table(s, 0.63, 1.32, 12.07, [2.6, 6.27, 1.9, 1.3], rows_o, size=9.5, row_h=0.43, bold_cols=(0,))
new_slides.append(s)

# ---- IP & DNS
s = content_slide("New Components ｜ IP & DNS Planning", "新增元件的 IP / FQDN 需求（實際值填在 Excel『IP & DNS 規劃』頁）", PLUM,
                  "依據：KB 440630；TechDocs《Deploy Components by Using VCF Installer…》《Upgrading to VCF Automation 9.1》；CXS Lab 實測")
rows_i = [["元件", "IP", "FQDN", "網段", "備註"]] + [list(r) for r in IPDNS]
table(s, 0.63, 1.32, 12.07, [3.1, 2.0, 2.2, 1.35, 3.42], rows_i, size=9, row_h=0.38, bold_cols=(0,))
new_slides.append(s)

# ---- references
s = content_slide("References ｜ Broadcom TechDocs & KB", "本計畫引用的官方依據（完整 URL 見 Excel『官方文件』頁）", PLUM)
left = [[(f"{cat}　", {"bold": True, "color": PLUM, "size": 8.5}), (t, {"size": 8.5})] for cat, t, u in DOCS]
tbox(s, 0.63, 1.3, 6.3, 5.4, [[("TechDocs（VCF 9.1 / 9.1.1）", {"bold": True, "color": NAVY, "size": 11})]] + left, size=8.5, space_after=2)
right = [[(f"KB {n}　", {"bold": True, "color": TEAL, "size": 8.5}), (t, {"size": 8.5})] for n, t in KBS]
tbox(s, 7.1, 1.3, 5.6, 5.4, [[("Knowledge Base", {"bold": True, "color": NAVY, "size": 11})]] + right, size=8.5, space_after=2)
new_slides.append(s)

# ---- big statement
s = prs.slides.add_slide(layout("Big Statement - Plum"))
for ph in s.placeholders:
    if ph.placeholder_format.idx == 12:
        ph.text = "先把管理網域做對，再一朵雲一朵雲接進來；每一步都有退路，PONR 前停下來簽核。"
        for r in ph.text_frame.paragraphs[0].runs:
            rPr = r._r.get_or_add_rPr(); ea = etree.SubElement(rPr, qn("a:ea")); ea.set("typeface", EA)
new_slides.append(s)

# ================================================================= cover text
cover = prs.slides[0]
for ph in cover.placeholders:
    idx = ph.placeholder_format.idx
    val = {0: "VCF 9.1.1 Converge & Import Plan", 10: "CTBC ｜ 執行計畫與每步 Checklist",
           11: "Kosten Yang", 12: "Broadcom CXS", 14: "October 2026"}.get(idx)
    if val is None: continue
    tf = ph.text_frame
    # keep first paragraph formatting
    p0 = tf.paragraphs[0]
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    for r in list(p0.runs)[1:]:
        r._r.getparent().remove(r._r)
    for br in p0._p.findall(qn("a:br")):
        p0._p.remove(br)
    if p0.runs:
        p0.runs[0].text = val
    else:
        p0.add_run().text = val
    rPr = p0.runs[0]._r.get_or_add_rPr()
    if rPr.find(qn("a:ea")) is None:
        ea = etree.SubElement(rPr, qn("a:ea")); ea.set("typeface", EA)

# ================================================================= reorder / drop template slides
sldIdLst = prs.slides._sldIdLst
all_ids = list(sldIdLst)
new_ids = [x for x in all_ids if x not in orig_ids]
keep = [COVER_ID] + new_ids + [THANKS_ID]
for sid in all_ids:
    if sid not in keep:
        prs.part.drop_rel(sid.rId)
    sldIdLst.remove(sid)
for sid in keep:
    sldIdLst.append(sid)

# strip speaker notes from the two reused template slides
for sl in (prs.slides[0], prs.slides[len(prs.slides) - 1]):
    if sl.has_notes_slide:
        sl.notes_slide.notes_text_frame.text = ""

import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
